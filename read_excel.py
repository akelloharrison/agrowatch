import openpyxl
import json

def read_excel_file(filename):
    print(f"\n{'='*60}")
    print(f"=== {filename} ===")
    print(f"{'='*60}")
    
    try:
        wb = openpyxl.load_workbook(filename)
        print(f"Sheet names: {wb.sheetnames}\n")
        
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            print(f"\n--- Sheet: {sheet_name} ---")
            print(f"Dimensions: {ws.dimensions}")
            print(f"\nData:\n")
            
            max_rows = min(50, ws.max_row)
            for i, row in enumerate(ws.iter_rows(min_row=1, max_row=max_rows, values_only=True), 1):
                print(f"{i:3d}: {row}")
    except Exception as e:
        print(f"Error: {e}")

read_excel_file('convertcsv (5).xlsx')
read_excel_file('convertcsv (7).xlsx')
